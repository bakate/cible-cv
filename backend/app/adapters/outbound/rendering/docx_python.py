"""Word (.docx) rendering with python-docx."""
import io
from datetime import datetime
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from app.domain.models import CvData, LetterData
from app.domain.rules import accent_of

_MUTED = RGBColor(0x55, 0x55, 0x55)


def _accent_rgb(hex_color: str) -> RGBColor:
    h = (hex_color or "#FF3E1A").lstrip("#")
    if len(h) != 6:
        h = "FF3E1A"
    return RGBColor.from_string(h.upper())


def _new_document(font_size: float) -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(font_size)
    return doc


def _heading(doc: Document, title: str, accent: RGBColor) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = accent


def _header(doc: Document, cv: CvData, accent: RGBColor) -> None:
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(2)
    nrun = name_p.add_run(cv.get("full_name", "") or "")
    nrun.bold = True
    nrun.font.size = Pt(22)
    if cv.get("headline"):
        head_p = doc.add_paragraph()
        head_p.paragraph_format.space_after = Pt(4)
        hr = head_p.add_run(cv["headline"])
        hr.bold = True
        hr.font.size = Pt(12)
        hr.font.color.rgb = accent
    contact = cv.get("contact") or {}
    parts = [contact[k] for k in ("email", "phone", "location", "linkedin", "github", "website") if contact.get(k)]
    if parts:
        cp = doc.add_paragraph(" · ".join(parts))
        cp.runs[0].font.size = Pt(9)
        cp.runs[0].font.color.rgb = _MUTED


def _experiences(doc: Document, experiences: List[Dict[str, Any]], accent: RGBColor) -> None:
    if not experiences:
        return
    _heading(doc, "Expériences professionnelles", accent)
    for e in experiences:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(e.get("title", "") or "")
        r.bold = True
        r.font.size = Pt(11)
        details = f"  ·  {e.get('company') or ''}"
        if e.get("location"):
            details += f" — {e['location']}"
        details += f"   |   {e.get('start', '')} – {e.get('end', '')}"
        sub = p.add_run(details)
        sub.font.size = Pt(10)
        sub.font.color.rgb = _MUTED
        for b in (e.get("bullets") or []):
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(0)


def _education(doc: Document, education: List[Dict[str, Any]], accent: RGBColor) -> None:
    if not education:
        return
    _heading(doc, "Formation", accent)
    for ed in education:
        p = doc.add_paragraph()
        r = p.add_run(ed.get("degree", "") or "")
        r.bold = True
        tail = f"  ·  {ed.get('school', '')}"
        dates = f"{ed.get('start', '')} – {ed.get('end', '')}"
        if dates.strip(" – "):
            tail += f"   |   {dates}"
        sub = p.add_run(tail)
        sub.font.size = Pt(10)
        sub.font.color.rgb = _MUTED
        if ed.get("details"):
            dp = doc.add_paragraph(ed["details"])
            dp.runs[0].italic = True
            dp.runs[0].font.size = Pt(10)


def _skills(doc: Document, cv: CvData, accent: RGBColor) -> None:
    if cv.get("skill_groups"):
        _heading(doc, "Compétences", accent)
        for g in cv["skill_groups"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{g.get('category', '')} : ")
            r.bold = True
            p.add_run(" · ".join(g.get("items") or []))
        return
    if cv.get("skills"):
        _heading(doc, "Compétences", accent)
        doc.add_paragraph(" · ".join(cv["skills"]))
    if cv.get("tools"):
        _heading(doc, "Outils", accent)
        doc.add_paragraph(" · ".join(cv["tools"]))


def _extras(doc: Document, cv: CvData, accent: RGBColor) -> None:
    if cv.get("languages"):
        _heading(doc, "Langues", accent)
        doc.add_paragraph(" · ".join(
            f"{lang.get('name', '')} ({lang.get('level', '')})" for lang in cv["languages"] if lang.get("name")
        ))
    if cv.get("certifications"):
        _heading(doc, "Certifications", accent)
        for cer in cv["certifications"]:
            line = cer.get("name", "") or ""
            if cer.get("issuer"):
                line += f" — {cer['issuer']}"
            if cer.get("year"):
                line += f" ({cer['year']})"
            doc.add_paragraph(line)
    if cv.get("interests"):
        _heading(doc, "Centres d'intérêt", accent)
        doc.add_paragraph(" · ".join(cv["interests"]))


def _save(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_cv_docx(cv: CvData) -> bytes:
    accent = _accent_rgb(accent_of(cv))
    doc = _new_document(10.5)
    _header(doc, cv, accent)
    if cv.get("summary"):
        _heading(doc, "Profil", accent)
        doc.add_paragraph(cv["summary"])
    _experiences(doc, cv.get("experiences") or [], accent)
    _education(doc, cv.get("education") or [], accent)
    _skills(doc, cv, accent)
    _extras(doc, cv, accent)
    return _save(doc)


def render_letter_docx(letter: LetterData, sender_cv: CvData, company: str) -> bytes:
    doc = _new_document(11)
    contact = (sender_cv or {}).get("contact") or {}
    sp = doc.add_paragraph()
    sname = sp.add_run(sender_cv.get("full_name", "") or "")
    sname.bold = True
    sub = [contact[k] for k in ("email", "phone", "location") if contact.get(k)]
    if sub:
        sp.add_run("\n" + "\n".join(sub))
    rp = doc.add_paragraph()
    rp.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    rr = rp.add_run(company or "")
    rr.bold = True
    rp.add_run(f"\n\n{datetime.now().strftime('%d/%m/%Y')}")
    if letter.get("subject"):
        op = doc.add_paragraph()
        r = op.add_run("Objet : ")
        r.bold = True
        op.add_run(letter["subject"])
    if letter.get("recipient"):
        doc.add_paragraph(letter["recipient"])
    for para in (letter.get("body") or "").split("\n"):
        doc.add_paragraph(para)
    return _save(doc)
