"""Cible CV - Backend FastAPI.

Generates a tailored CV and motivation letter from a profile and a job offer
using Claude Sonnet 4.5 via the Emergent integration library.
"""
import os
import io
import re
import json
import uuid
import logging
from typing import List, Optional, Any, Dict

import requests
from bs4 import BeautifulSoup
import html2text
from pypdf import PdfReader
from docx import Document

from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException, Form, Depends, Request, Cookie, Header, Response as FResponse
from fastapi.responses import JSONResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, ConfigDict
from pathlib import Path
from datetime import datetime, timezone, timedelta

import requests
from bs4 import BeautifulSoup
import html2text
from pypdf import PdfReader
from docx import Document

from emergentintegrations.llm.chat import LlmChat, UserMessage

from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from fastapi.responses import Response

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.platypus.flowables import HRFlowable
from reportlab.lib.enums import TA_RIGHT

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger("ciblecv")

# ---- MongoDB ----
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]
mongo_client = AsyncIOMotorClient(MONGO_URL)
db = mongo_client[DB_NAME]

EMERGENT_LLM_KEY = os.environ.get("EMERGENT_LLM_KEY", "")
ANTHROPIC_MODEL = "claude-sonnet-4-5-20250929"
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "").strip().lower()
EMERGENT_AUTH_BASE = "https://demobackend.emergentagent.com/auth/v1/env"

app = FastAPI(title="Cible CV API")
api = APIRouter(prefix="/api")


# ---------- Auth models ----------
class User(BaseModel):
    user_id: str
    email: str
    name: str
    picture: Optional[str] = None
    role: str = "user"
    created_at: str


class AuthMeResponse(BaseModel):
    user_id: str
    email: str
    name: str
    picture: Optional[str] = None
    role: str


# ---------- Models ----------
class GenerationCreate(BaseModel):
    profile_text: str
    profile_meta: Optional[Dict[str, Any]] = None
    job_text: str
    job_meta: Optional[Dict[str, Any]] = None
    template: str = "corporate"  # "corporate" | "startup"
    photo_data_url: Optional[str] = None


class GenerationUpdate(BaseModel):
    cv: Optional[Dict[str, Any]] = None
    letter: Optional[Dict[str, Any]] = None
    template: Optional[str] = None
    photo_data_url: Optional[str] = None


class GenerationRecord(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    title: str
    company: str
    position: str
    template: str
    profile_text: str
    profile_meta: Optional[Dict[str, Any]] = None
    job_text: str
    job_meta: Optional[Dict[str, Any]] = None
    cv: Dict[str, Any]
    letter: Dict[str, Any]
    adaptations: Dict[str, Any]
    photo_data_url: Optional[str] = None
    created_at: str


# ---------- Utilities ----------
def extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as e:
            logger.warning("pdf page extract failed: %s", e)
    return "\n".join(parts).strip()


def extract_docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip()


def fetch_url_text(url: str) -> str:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36"
        ),
        "Accept-Language": "fr-FR,fr;q=0.9,en;q=0.8",
    }
    r = requests.get(url, headers=headers, timeout=15)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "lxml")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header"]):
        tag.decompose()
    # Focus on main / article when available
    main = soup.find("main") or soup.find("article") or soup.body or soup
    h = html2text.HTML2Text()
    h.ignore_links = True
    h.ignore_images = True
    h.body_width = 0
    text = h.handle(str(main))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:20000]


def _extract_json(raw: str) -> Dict[str, Any]:
    """Pull the first JSON object/array out of an LLM response."""
    fenced = re.search(r"```(?:json)?\s*(.*?)```", raw, re.DOTALL)
    candidate = fenced.group(1) if fenced else raw
    start = candidate.find("{")
    end = candidate.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("No JSON object found in LLM response")
    return json.loads(candidate[start : end + 1])


async def claude_chat(system_prompt: str, user_text: str) -> str:
    if not EMERGENT_LLM_KEY:
        raise HTTPException(status_code=500, detail="EMERGENT_LLM_KEY non configurée")
    chat = (
        LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"cible-cv-{uuid.uuid4()}",
            system_message=system_prompt,
        )
        .with_model("anthropic", ANTHROPIC_MODEL)
        .with_params(max_tokens=8000)
    )
    response = await chat.send_message(UserMessage(text=user_text))
    return response if isinstance(response, str) else str(response)


# ---------- Routes ----------
@api.get("/")
async def root():
    return {"app": "Cible CV", "status": "ok"}


# ---------- Auth helpers ----------
SESSION_COOKIE = "session_token"
SESSION_TTL_DAYS = 7


def _extract_session_token(
    cookie_token: Optional[str],
    auth_header: Optional[str],
) -> Optional[str]:
    if cookie_token:
        return cookie_token
    if auth_header and auth_header.lower().startswith("bearer "):
        return auth_header.split(" ", 1)[1].strip()
    return None


async def _user_from_session(token: str) -> Optional[Dict[str, Any]]:
    session = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if not session:
        return None
    expires_at = session.get("expires_at")
    if isinstance(expires_at, str):
        try:
            expires_at = datetime.fromisoformat(expires_at)
        except ValueError:
            return None
    if expires_at and expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=timezone.utc)
    if expires_at and expires_at < datetime.now(timezone.utc):
        return None
    user = await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0})
    return user


async def get_current_user(
    session_token: Optional[str] = Cookie(default=None),
    authorization: Optional[str] = Header(default=None),
) -> Dict[str, Any]:
    token = _extract_session_token(session_token, authorization)
    if not token:
        raise HTTPException(status_code=401, detail="Non authentifié")
    user = await _user_from_session(token)
    if not user:
        raise HTTPException(status_code=401, detail="Session invalide ou expirée")
    return user


async def require_admin(user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Réservé aux admins")
    return user


def _set_session_cookie(response: FResponse, token: str) -> None:
    response.set_cookie(
        key=SESSION_COOKIE,
        value=token,
        max_age=SESSION_TTL_DAYS * 24 * 3600,
        httponly=True,
        secure=True,
        samesite="none",
        path="/",
    )


async def _ensure_user_and_session(profile: Dict[str, Any]) -> Dict[str, Any]:
    """Upsert user from Emergent profile, decide role, store session, return user."""
    email = (profile.get("email") or "").strip().lower()
    if not email:
        raise HTTPException(status_code=400, detail="Email manquant dans le profil OAuth")

    existing = await db.users.find_one({"email": email}, {"_id": 0})
    now_iso = datetime.now(timezone.utc).isoformat()

    if existing:
        user = existing
        update = {
            "name": profile.get("name") or existing.get("name"),
            "picture": profile.get("picture") or existing.get("picture"),
        }
        await db.users.update_one({"user_id": user["user_id"]}, {"$set": update})
        user.update(update)
    else:
        admin_count = await db.users.count_documents({"role": "admin"})
        role = "admin" if (
            (ADMIN_EMAIL and email == ADMIN_EMAIL) or admin_count == 0
        ) else "user"
        user = {
            "user_id": f"user_{uuid.uuid4().hex[:12]}",
            "email": email,
            "name": profile.get("name") or email.split("@")[0],
            "picture": profile.get("picture"),
            "role": role,
            "created_at": now_iso,
        }
        await db.users.insert_one(user)
        if role == "admin":
            # Migrate orphan data (no user_id) to this first admin.
            await db.generations.update_many(
                {"user_id": {"$exists": False}},
                {"$set": {"user_id": user["user_id"]}},
            )
            await db.base_profile.update_many(
                {"user_id": {"$exists": False}},
                {"$set": {"user_id": user["user_id"]}},
            )

    session_token = profile.get("session_token") or uuid.uuid4().hex
    expires_at = datetime.now(timezone.utc) + timedelta(days=SESSION_TTL_DAYS)
    await db.user_sessions.insert_one(
        {
            "user_id": user["user_id"],
            "session_token": session_token,
            "expires_at": expires_at,
            "created_at": datetime.now(timezone.utc),
        }
    )
    user["session_token"] = session_token
    return user


# ---------- Auth endpoints ----------
@api.post("/auth/process-session")
async def auth_process_session(
    response: FResponse,
    x_session_id: Optional[str] = Header(default=None, alias="X-Session-ID"),
):
    if not x_session_id:
        raise HTTPException(status_code=400, detail="X-Session-ID requis")
    try:
        r = requests.get(
            f"{EMERGENT_AUTH_BASE}/oauth/session-data",
            headers={"X-Session-ID": x_session_id},
            timeout=10,
        )
        r.raise_for_status()
        profile = r.json()
    except Exception as e:
        logger.exception("Emergent session-data fetch failed")
        raise HTTPException(status_code=502, detail=f"OAuth fetch failed: {e}")

    user = await _ensure_user_and_session(profile)
    _set_session_cookie(response, user["session_token"])
    return {
        "user_id": user["user_id"],
        "email": user["email"],
        "name": user["name"],
        "picture": user.get("picture"),
        "role": user["role"],
    }


@api.get("/auth/me", response_model=AuthMeResponse)
async def auth_me(user: Dict[str, Any] = Depends(get_current_user)):
    return AuthMeResponse(
        user_id=user["user_id"],
        email=user["email"],
        name=user["name"],
        picture=user.get("picture"),
        role=user.get("role", "user"),
    )


@api.post("/auth/logout")
async def auth_logout(
    response: FResponse,
    session_token: Optional[str] = Cookie(default=None),
):
    if session_token:
        await db.user_sessions.delete_one({"session_token": session_token})
    response.delete_cookie(SESSION_COOKIE, path="/", samesite="none", secure=True)
    return {"ok": True}


@api.post("/parse/pdf")
async def parse_pdf(file: UploadFile = File(...), _user: Dict[str, Any] = Depends(get_current_user)):
    content = await file.read()
    name = (file.filename or "").lower()
    text = ""
    try:
        if name.endswith(".pdf"):
            text = extract_pdf_text(content)
        elif name.endswith(".docx"):
            text = extract_docx_text(content)
        elif name.endswith(".txt"):
            text = content.decode("utf-8", errors="ignore")
        else:
            raise HTTPException(status_code=400, detail="Format non supporté (PDF, DOCX, TXT)")
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("parse failed")
        raise HTTPException(status_code=400, detail=f"Lecture du fichier impossible: {e}")
    if not text:
        raise HTTPException(status_code=400, detail="Le document est vide ou non lisible.")
    return {"text": text, "filename": file.filename}


@api.post("/parse/url")
async def parse_url(payload: Dict[str, str], _user: Dict[str, Any] = Depends(get_current_user)):
    url = (payload.get("url") or "").strip()
    if not url:
        raise HTTPException(status_code=400, detail="URL requise")
    text = ""
    try:
        text = fetch_url_text(url)
    except Exception as e:
        logger.warning("URL fetch failed for %s: %s", url, e)
        raise HTTPException(
            status_code=400,
            detail="Impossible de récupérer cette URL automatiquement. Collez le contenu à la main.",
        )
    if not text or len(text) < 80:
        raise HTTPException(
            status_code=400,
            detail="Contenu insuffisant à cette URL. Collez le texte de l'annonce manuellement.",
        )
    return {"text": text, "url": url}


SYSTEM_PROMPT = """Tu es un expert RH et coach carrière francophone. Tu produis des CV et lettres de motivation en français, percutants, professionnels et rigoureusement honnêtes.

RÈGLES CRITIQUES:
1. N'invente JAMAIS d'expérience, de diplôme, de compétence ou de chiffre.
2. Si une information manque, reste sobre, ne comble pas avec du fictif.
3. Adapte le vocabulaire au champ lexical de l'annonce, sans mentir.
4. Réorganise et reformule pour mettre en avant ce qui matche l'offre.
5. La lettre doit être naturelle, crédible, sans formules clichées creuses.
6. Réponds STRICTEMENT en JSON valide, sans texte avant/après, sans markdown.
"""


GENERATION_TEMPLATE = """Voici le PROFIL CANDIDAT (extrait brut du CV / LinkedIn / saisie manuelle):
---
{profile}
---

Voici l'OFFRE D'EMPLOI cible:
---
{job}
---

Génère un objet JSON structuré exactement comme suit:

{{
  "cv": {{
    "full_name": "string",
    "headline": "intitulé court adapté au poste",
    "contact": {{
      "email": "string ou null",
      "phone": "string ou null",
      "location": "string ou null",
      "linkedin": "string ou null",
      "github": "string ou null",
      "website": "string ou null"
    }},
    "summary": "résumé pro 3-4 lignes orienté annonce",
    "skill_groups": [
      {{"category": "ex: Frontend / Backend / DevOps / Sécurité / IA / Data / Méthodologies / Outils", "items": ["compétence ou outil concret", "..."]}}
    ],
    "skills": ["compétence", "..."],
    "languages": [{{"name": "Français", "level": "Natif"}}],
    "experiences": [
      {{
        "title": "poste",
        "company": "entreprise",
        "location": "ville",
        "start": "MM/YYYY",
        "end": "MM/YYYY ou Présent",
        "bullets": ["réalisation orientée résultats", "..."]
      }}
    ],
    "education": [
      {{"degree": "diplôme", "school": "école", "start": "YYYY", "end": "YYYY", "details": "string ou null"}}
    ],
    "certifications": [{{"name": "string", "issuer": "string", "year": "YYYY"}}],
    "tools": ["outil", "..."],
    "interests": ["centre d'intérêt", "..."]
  }},
  "letter": {{
    "recipient": "Madame, Monsieur,",
    "subject": "Candidature au poste de ...",
    "body": "Lettre complète en 3-4 paragraphes naturels, en français. Personnalisée à l'entreprise et au poste. Pas de répétition mot pour mot du CV. Termine par une formule de politesse simple."
  }},
  "adaptations": {{
    "position": "poste détecté",
    "company": "entreprise détectée",
    "keywords_matched": ["mot-clé présent chez le candidat ET dans l'offre"],
    "keywords_added": ["mot-clé reformulé / mis en avant"],
    "experiences_highlighted": ["expérience mise en avant et pourquoi"],
    "match_score": 0
  }}
}}

Le match_score est un entier de 0 à 100 reflétant la correspondance globale.
"skill_groups" doit être bien rempli : 5 à 8 familles thématiques (ex. Frontend, Backend, DevOps & Cloud, Sécurité, IA, Data, Méthodologies, Outils) avec les compétences/outils du candidat ventilés dedans. La liste "skills" plate reste utile pour les ATS, les deux doivent contenir les mêmes éléments.
RAPPEL: Reste fidèle au parcours réel du candidat. Sors UNIQUEMENT le JSON.
"""


@api.post("/generate")
async def generate(req: GenerationCreate, user: Dict[str, Any] = Depends(get_current_user)):
    if not req.profile_text or len(req.profile_text) < 50:
        raise HTTPException(status_code=400, detail="Profil insuffisant pour générer.")
    if not req.job_text or len(req.job_text) < 50:
        raise HTTPException(status_code=400, detail="Offre d'emploi insuffisante.")
    prompt = GENERATION_TEMPLATE.format(
        profile=req.profile_text[:14000],
        job=req.job_text[:8000],
    )
    data: Dict[str, Any] = {}
    try:
        raw = await claude_chat(SYSTEM_PROMPT, prompt)
        data = _extract_json(raw)
    except Exception as e:
        logger.exception("LLM generation failed")
        raise HTTPException(status_code=502, detail=f"Échec de la génération: {e}")

    cv = data.get("cv") or {}
    letter = data.get("letter") or {}
    adaptations = data.get("adaptations") or {}

    rec_id = str(uuid.uuid4())
    now_iso = datetime.now(timezone.utc).isoformat()
    company = adaptations.get("company") or (req.job_meta or {}).get("company") or "Entreprise"
    position = adaptations.get("position") or "Poste"
    title = f"{position} — {company}"

    doc = {
        "id": rec_id,
        "user_id": user["user_id"],
        "title": title,
        "company": company,
        "position": position,
        "template": req.template,
        "profile_text": req.profile_text,
        "profile_meta": req.profile_meta or {},
        "job_text": req.job_text,
        "job_meta": req.job_meta or {},
        "cv": cv,
        "letter": letter,
        "adaptations": adaptations,
        "photo_data_url": req.photo_data_url,
        "created_at": now_iso,
    }
    await db.generations.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.get("/generations")
async def list_generations(user: Dict[str, Any] = Depends(get_current_user)):
    query = {} if user.get("role") == "admin" else {"user_id": user["user_id"]}
    cursor = db.generations.find(
        query,
        {
            "_id": 0,
            "id": 1,
            "title": 1,
            "company": 1,
            "position": 1,
            "template": 1,
            "created_at": 1,
            "adaptations": 1,
            "user_id": 1,
        },
    ).sort("created_at", -1).limit(200)
    items = await cursor.to_list(200)
    for it in items:
        adp = it.pop("adaptations", {}) or {}
        it["match_score"] = adp.get("match_score", 0)
    return items


@api.get("/generations/{gen_id}")
async def get_generation(gen_id: str, user: Dict[str, Any] = Depends(get_current_user)):
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    doc = await db.generations.find_one(query, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    return doc


@api.put("/generations/{gen_id}")
async def update_generation(
    gen_id: str,
    update: GenerationUpdate,
    user: Dict[str, Any] = Depends(get_current_user),
):
    patch = {k: v for k, v in update.model_dump(exclude_none=True).items()}
    if not patch:
        raise HTTPException(status_code=400, detail="Rien à mettre à jour")
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    result = await db.generations.update_one(query, {"$set": patch})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    doc = await db.generations.find_one({"id": gen_id}, {"_id": 0})
    return doc


@api.delete("/generations/{gen_id}")
async def delete_generation(gen_id: str, user: Dict[str, Any] = Depends(get_current_user)):
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    result = await db.generations.delete_one(query)
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    return {"ok": True}


# ---- DOCX export ----
def _accent_rgb(hex_color: str) -> RGBColor:
    h = (hex_color or "#FF3E1A").lstrip("#")
    if len(h) != 6:
        h = "FF3E1A"
    return RGBColor.from_string(h.upper())


def _set_doc_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)


def _section_heading(doc: Document, title: str, accent: RGBColor) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = accent


_MUTED = RGBColor(0x55, 0x55, 0x55)


def _docx_header(doc: Document, cv: Dict[str, Any], accent: RGBColor) -> None:
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(2)
    nrun = name_p.add_run(cv.get("full_name", "") or "")
    nrun.bold = True
    nrun.font.size = Pt(22)

    if cv.get("headline"):
        head_p = doc.add_paragraph()
        head_p.paragraph_format.space_after = Pt(4)
        hr = head_p.add_run(cv["headline"])
        hr.bold = True
        hr.font.size = Pt(12)
        hr.font.color.rgb = accent

    contact = cv.get("contact") or {}
    parts = [contact[k] for k in ("email", "phone", "location", "linkedin", "github", "website") if contact.get(k)]
    if parts:
        cp = doc.add_paragraph(" · ".join(parts))
        cp.runs[0].font.size = Pt(9)
        cp.runs[0].font.color.rgb = _MUTED


def _docx_experiences(doc: Document, experiences: List[Dict[str, Any]], accent: RGBColor) -> None:
    if not experiences:
        return
    _section_heading(doc, "Expériences professionnelles", accent)
    for e in experiences:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(e.get("title", "") or "")
        r.bold = True
        r.font.size = Pt(11)
        loc = e.get("location") or ""
        details = f"  ·  {e.get('company') or ''}"
        if loc:
            details += f" — {loc}"
        details += f"   |   {e.get('start', '')} – {e.get('end', '')}"
        sub = p.add_run(details)
        sub.font.size = Pt(10)
        sub.font.color.rgb = _MUTED
        for b in (e.get("bullets") or []):
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(0)


def _docx_education(doc: Document, education: List[Dict[str, Any]], accent: RGBColor) -> None:
    if not education:
        return
    _section_heading(doc, "Formation", accent)
    for ed in education:
        p = doc.add_paragraph()
        r = p.add_run(ed.get("degree", "") or "")
        r.bold = True
        tail = f"  ·  {ed.get('school', '')}"
        dates = f"{ed.get('start', '')} – {ed.get('end', '')}"
        if dates.strip(" – "):
            tail += f"   |   {dates}"
        sub = p.add_run(tail)
        sub.font.size = Pt(10)
        sub.font.color.rgb = _MUTED
        if ed.get("details"):
            dp = doc.add_paragraph(ed["details"])
            dp.runs[0].italic = True
            dp.runs[0].font.size = Pt(10)


def _docx_skills(doc: Document, cv: Dict[str, Any], accent: RGBColor) -> None:
    if cv.get("skill_groups"):
        _section_heading(doc, "Compétences", accent)
        for g in cv["skill_groups"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{g.get('category', '')} : ")
            r.bold = True
            p.add_run(" · ".join(g.get("items") or []))
        return
    if cv.get("skills"):
        _section_heading(doc, "Compétences", accent)
        doc.add_paragraph(" · ".join(cv["skills"]))
    if cv.get("tools"):
        _section_heading(doc, "Outils", accent)
        doc.add_paragraph(" · ".join(cv["tools"]))


def _docx_extras(doc: Document, cv: Dict[str, Any], accent: RGBColor) -> None:
    if cv.get("languages"):
        _section_heading(doc, "Langues", accent)
        doc.add_paragraph(
            " · ".join(
                f"{lang.get('name', '')} ({lang.get('level', '')})"
                for lang in cv["languages"]
                if lang.get("name")
            )
        )
    if cv.get("certifications"):
        _section_heading(doc, "Certifications", accent)
        for cer in cv["certifications"]:
            line = cer.get("name", "") or ""
            if cer.get("issuer"):
                line += f" — {cer['issuer']}"
            if cer.get("year"):
                line += f" ({cer['year']})"
            doc.add_paragraph(line)
    if cv.get("interests"):
        _section_heading(doc, "Centres d'intérêt", accent)
        doc.add_paragraph(" · ".join(cv["interests"]))


def build_cv_docx(cv: Dict[str, Any]) -> bytes:
    accent_hex = ((cv.get("theme") or {}).get("accent")) or "#FF3E1A"
    accent = _accent_rgb(accent_hex)
    doc = Document()
    _set_doc_margins(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    _docx_header(doc, cv, accent)

    if cv.get("summary"):
        _section_heading(doc, "Profil", accent)
        doc.add_paragraph(cv["summary"])

    _docx_experiences(doc, cv.get("experiences") or [], accent)
    _docx_education(doc, cv.get("education") or [], accent)
    _docx_skills(doc, cv, accent)
    _docx_extras(doc, cv, accent)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_letter_docx(letter: Dict[str, Any], sender_cv: Dict[str, Any], company: str) -> bytes:
    doc = Document()
    _set_doc_margins(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Sender block
    sender_contact = (sender_cv or {}).get("contact") or {}
    sp = doc.add_paragraph()
    sname = sp.add_run(sender_cv.get("full_name", "") or "")
    sname.bold = True
    sub = []
    for k in ("email", "phone", "location"):
        v = sender_contact.get(k)
        if v:
            sub.append(v)
    if sub:
        sp.add_run("\n" + "\n".join(sub))

    # Recipient + date
    rp = doc.add_paragraph()
    rp.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    rr = rp.add_run(company or "")
    rr.bold = True
    today = datetime.now().strftime("%d/%m/%Y")
    rp.add_run(f"\n\n{today}")

    if letter.get("subject"):
        sp = doc.add_paragraph()
        r = sp.add_run("Objet : ")
        r.bold = True
        sp.add_run(letter["subject"])

    if letter.get("recipient"):
        doc.add_paragraph(letter["recipient"])

    for para in (letter.get("body") or "").split("\n"):
        doc.add_paragraph(para)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", name or "candidat").strip("_") or "candidat"


# ---------- Text-based PDF (ATS-friendly) ----------
def _esc(s: Any) -> str:
    if s is None:
        return ""
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _pdf_styles(accent_hex: str) -> Dict[str, ParagraphStyle]:
    accent = colors.HexColor(accent_hex)
    muted = colors.HexColor("#555555")
    return {
        "name": ParagraphStyle("Name", fontSize=22, leading=24, fontName="Helvetica-Bold", spaceAfter=2),
        "headline": ParagraphStyle("Headline", fontSize=12, leading=15, fontName="Helvetica-Bold", textColor=accent, spaceAfter=4),
        "contact": ParagraphStyle("Contact", fontSize=9, leading=11, fontName="Helvetica", textColor=muted),
        "section": ParagraphStyle("Section", fontSize=10.5, leading=13, fontName="Helvetica-Bold", textColor=accent, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("Body", fontSize=10, leading=13, fontName="Helvetica"),
        "exp_title": ParagraphStyle("ExpTitle", fontSize=11, leading=13, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=0),
        "exp_meta": ParagraphStyle("ExpMeta", fontSize=9.5, leading=12, fontName="Helvetica-Oblique", textColor=muted, spaceAfter=2),
        "bullet": ParagraphStyle("Bullet", fontSize=10, leading=13, fontName="Helvetica", leftIndent=12),
        "right": ParagraphStyle("Right", fontSize=11, leading=14, fontName="Helvetica", alignment=TA_RIGHT),
        "letter_body": ParagraphStyle("LetterBody", fontSize=11, leading=15, fontName="Helvetica", spaceAfter=6),
        "letter_bold": ParagraphStyle("LetterBold", fontSize=11, leading=15, fontName="Helvetica-Bold"),
    }


def _pdf_header(flow: List, cv: Dict[str, Any], styles: Dict[str, ParagraphStyle]) -> None:
    flow.append(Paragraph(_esc(cv.get("full_name", "")), styles["name"]))
    if cv.get("headline"):
        flow.append(Paragraph(_esc(cv["headline"]), styles["headline"]))
    contact = cv.get("contact") or {}
    parts = [contact[k] for k in ("email", "phone", "location", "linkedin", "github", "website") if contact.get(k)]
    if parts:
        flow.append(Paragraph(" &nbsp;·&nbsp; ".join(_esc(p) for p in parts), styles["contact"]))
    flow.append(HRFlowable(width="100%", thickness=1.2, color=colors.black, spaceBefore=4, spaceAfter=4))


def _pdf_experiences(flow: List, experiences: List[Dict[str, Any]], styles: Dict[str, ParagraphStyle]) -> None:
    if not experiences:
        return
    flow.append(Paragraph("EXPÉRIENCES PROFESSIONNELLES", styles["section"]))
    for e in experiences:
        title = _esc(e.get("title", ""))
        company = _esc(e.get("company", ""))
        flow.append(Paragraph(f"<b>{title}</b> · {company}", styles["exp_title"]))
        loc = _esc(e.get("location", ""))
        dates = f"{_esc(e.get('start', ''))} – {_esc(e.get('end', ''))}"
        meta = f"{loc} &nbsp;|&nbsp; {dates}" if loc else dates
        flow.append(Paragraph(meta, styles["exp_meta"]))
        for b in (e.get("bullets") or []):
            flow.append(Paragraph(f"•&nbsp; {_esc(b)}", styles["bullet"]))
        flow.append(Spacer(1, 4))


def _pdf_education(flow: List, education: List[Dict[str, Any]], styles: Dict[str, ParagraphStyle]) -> None:
    if not education:
        return
    flow.append(Paragraph("FORMATION", styles["section"]))
    for ed in education:
        flow.append(Paragraph(f"<b>{_esc(ed.get('degree', ''))}</b> · {_esc(ed.get('school', ''))}", styles["exp_title"]))
        dates = f"{_esc(ed.get('start', ''))} – {_esc(ed.get('end', ''))}"
        if ed.get("details"):
            dates += f" — {_esc(ed['details'])}"
        flow.append(Paragraph(dates, styles["exp_meta"]))


def _pdf_skills(flow: List, cv: Dict[str, Any], styles: Dict[str, ParagraphStyle]) -> None:
    if cv.get("skill_groups"):
        flow.append(Paragraph("COMPÉTENCES", styles["section"]))
        for g in cv["skill_groups"]:
            cat = _esc(g.get("category", ""))
            items = " · ".join(_esc(i) for i in (g.get("items") or []))
            flow.append(Paragraph(f"<b>{cat} :</b> {items}", styles["body"]))
            flow.append(Spacer(1, 2))
        return
    if cv.get("skills"):
        flow.append(Paragraph("COMPÉTENCES", styles["section"]))
        flow.append(Paragraph(" · ".join(_esc(s) for s in cv["skills"]), styles["body"]))
    if cv.get("tools"):
        flow.append(Paragraph("OUTILS", styles["section"]))
        flow.append(Paragraph(" · ".join(_esc(s) for s in cv["tools"]), styles["body"]))


def _pdf_extras(flow: List, cv: Dict[str, Any], styles: Dict[str, ParagraphStyle]) -> None:
    if cv.get("languages"):
        flow.append(Paragraph("LANGUES", styles["section"]))
        langs = " · ".join(
            f"{_esc(lang.get('name', ''))} ({_esc(lang.get('level', ''))})"
            for lang in cv["languages"]
            if lang.get("name")
        )
        flow.append(Paragraph(langs, styles["body"]))
    if cv.get("certifications"):
        flow.append(Paragraph("CERTIFICATIONS", styles["section"]))
        for cer in cv["certifications"]:
            line = _esc(cer.get("name", ""))
            if cer.get("issuer"):
                line += f" — {_esc(cer['issuer'])}"
            if cer.get("year"):
                line += f" ({_esc(cer['year'])})"
            flow.append(Paragraph(line, styles["body"]))
    if cv.get("interests"):
        flow.append(Paragraph("CENTRES D'INTÉRÊT", styles["section"]))
        flow.append(Paragraph(" · ".join(_esc(i) for i in cv["interests"]), styles["body"]))


def build_cv_pdf_text(cv: Dict[str, Any]) -> bytes:
    accent_hex = ((cv.get("theme") or {}).get("accent")) or "#FF3E1A"
    styles = _pdf_styles(accent_hex)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title=cv.get("full_name", "CV"),
    )
    flow: List = []
    _pdf_header(flow, cv, styles)
    if cv.get("summary"):
        flow.append(Paragraph("PROFIL", styles["section"]))
        flow.append(Paragraph(_esc(cv["summary"]), styles["body"]))
    _pdf_experiences(flow, cv.get("experiences") or [], styles)
    _pdf_education(flow, cv.get("education") or [], styles)
    _pdf_skills(flow, cv, styles)
    _pdf_extras(flow, cv, styles)
    doc.build(flow)
    return buf.getvalue()


def build_letter_pdf_text(letter: Dict[str, Any], sender_cv: Dict[str, Any], company: str) -> bytes:
    accent_hex = ((sender_cv.get("theme") or {}).get("accent")) or "#FF3E1A"
    styles = _pdf_styles(accent_hex)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=22 * mm, rightMargin=22 * mm,
        topMargin=22 * mm, bottomMargin=22 * mm,
        title="Lettre de motivation",
    )
    flow: List = []
    sender_contact = (sender_cv or {}).get("contact") or {}
    flow.append(Paragraph(f"<b>{_esc(sender_cv.get('full_name', ''))}</b>", styles["letter_body"]))
    for k in ("email", "phone", "location"):
        v = sender_contact.get(k)
        if v:
            flow.append(Paragraph(_esc(v), styles["letter_body"]))
    flow.append(Spacer(1, 10))
    flow.append(Paragraph(f"<b>{_esc(company)}</b>", styles["right"]))
    today = datetime.now().strftime("%d/%m/%Y")
    flow.append(Paragraph(today, styles["right"]))
    flow.append(Spacer(1, 16))
    if letter.get("subject"):
        flow.append(Paragraph(f"<b>Objet :</b> {_esc(letter['subject'])}", styles["letter_body"]))
    if letter.get("recipient"):
        flow.append(Paragraph(_esc(letter["recipient"]), styles["letter_body"]))
    for para in (letter.get("body") or "").split("\n"):
        if para.strip():
            flow.append(Paragraph(_esc(para), styles["letter_body"]))
        else:
            flow.append(Spacer(1, 6))
    doc.build(flow)
    return buf.getvalue()


@api.get("/generations/{gen_id}/export/cv.docx")
async def export_cv_docx(gen_id: str, user: Dict[str, Any] = Depends(get_current_user)):
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    doc = await db.generations.find_one(query, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    blob = build_cv_docx(doc.get("cv") or {})
    fn = f"CV-{_safe_name(doc.get('cv', {}).get('full_name', ''))}-{_safe_name(doc.get('company', ''))}.docx"
    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


@api.get("/generations/{gen_id}/export/cv.pdf")
async def export_cv_pdf(gen_id: str, user: Dict[str, Any] = Depends(get_current_user)):
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    doc = await db.generations.find_one(query, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    blob = build_cv_pdf_text(doc.get("cv") or {})
    fn = f"CV-{_safe_name(doc.get('cv', {}).get('full_name', ''))}-{_safe_name(doc.get('company', ''))}.pdf"
    return Response(
        content=blob,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


@api.get("/generations/{gen_id}/export/letter.pdf")
async def export_letter_pdf(gen_id: str, user: Dict[str, Any] = Depends(get_current_user)):
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    doc = await db.generations.find_one(query, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    blob = build_letter_pdf_text(
        doc.get("letter") or {},
        doc.get("cv") or {},
        doc.get("company") or "",
    )
    fn = f"Lettre-{_safe_name(doc.get('cv', {}).get('full_name', ''))}-{_safe_name(doc.get('company', ''))}.pdf"
    return Response(
        content=blob,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


@api.get("/generations/{gen_id}/export/letter.docx")
async def export_letter_docx(gen_id: str, user: Dict[str, Any] = Depends(get_current_user)):
    query = {"id": gen_id}
    if user.get("role") != "admin":
        query["user_id"] = user["user_id"]
    doc = await db.generations.find_one(query, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Génération introuvable")
    blob = build_letter_docx(
        doc.get("letter") or {},
        doc.get("cv") or {},
        doc.get("company") or "",
    )
    fn = f"Lettre-{_safe_name(doc.get('cv', {}).get('full_name', ''))}-{_safe_name(doc.get('company', ''))}.docx"
    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


# ---- Base profile ("CV de base") ----
@api.get("/profile/base")
async def get_base_profile(user: Dict[str, Any] = Depends(get_current_user)):
    doc = await db.base_profile.find_one({"user_id": user["user_id"]}, {"_id": 0})
    if not doc:
        return {"exists": False}
    return {"exists": True, **doc}


@api.put("/profile/base")
async def put_base_profile(
    payload: Dict[str, Any],
    user: Dict[str, Any] = Depends(get_current_user),
):
    doc = {
        "user_id": user["user_id"],
        "profile_text": payload.get("profile_text", ""),
        "cv": payload.get("cv") or {},
        "photo_data_url": payload.get("photo_data_url"),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.base_profile.update_one(
        {"user_id": user["user_id"]}, {"$set": doc}, upsert=True
    )
    return {"exists": True, **doc}


@api.delete("/profile/base")
async def delete_base_profile(user: Dict[str, Any] = Depends(get_current_user)):
    await db.base_profile.delete_one({"user_id": user["user_id"]})
    return {"ok": True}


@api.post("/regroup-skills")
async def regroup_skills(
    payload: Dict[str, Any],
    _user: Dict[str, Any] = Depends(get_current_user),
):
    skills = payload.get("skills") or []
    tools = payload.get("tools") or []
    if not skills and not tools:
        raise HTTPException(status_code=400, detail="Aucune compétence à regrouper")
    prompt = (
        "Voici une liste de compétences et d'outils techniques d'un candidat. Regroupe-les "
        "en 5 à 8 familles thématiques pertinentes (ex. Frontend, Backend, DevOps & Cloud, "
        "Sécurité, IA, Data, Méthodologies, Outils). Garde TOUS les éléments — n'en perds "
        "aucun. Réponds STRICTEMENT en JSON valide:\n\n"
        '{"skill_groups": [{"category": "string", "items": ["string", ...]}]}'
        f"\n\nCompétences:\n{json.dumps(skills, ensure_ascii=False)}"
        f"\n\nOutils:\n{json.dumps(tools, ensure_ascii=False)}"
    )
    data: Dict[str, Any] = {}
    try:
        raw = await claude_chat(SYSTEM_PROMPT, prompt)
        data = _extract_json(raw)
    except Exception as e:
        logger.exception("regroup failed")
        raise HTTPException(status_code=502, detail=f"Regroupement échoué: {e}")
    return data


app.include_router(api)
_cors_origins = os.environ.get("CORS_ORIGINS", "").strip()
if _cors_origins and _cors_origins != "*":
    app.add_middleware(
        CORSMiddleware,
        allow_credentials=True,
        allow_origins=[o.strip() for o in _cors_origins.split(",") if o.strip()],
        allow_methods=["*"],
        allow_headers=["*"],
    )
else:
    # Permissive but cookie-friendly: regex matches any origin while still allowing credentials
    app.add_middleware(
        CORSMiddleware,
        allow_credentials=True,
        allow_origin_regex=".*",
        allow_methods=["*"],
        allow_headers=["*"],
    )


@app.on_event("shutdown")
async def shutdown():
    mongo_client.close()
